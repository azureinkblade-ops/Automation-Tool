"""Build Fantasy_Author_Image_Pack_100.docx from the fantasy_pack manifest + images.
Layout: 4 images per page (2x2 grid) with captions; a clickable reference table
(100 rows) whose caption links to each image's bookmark.
Mirrors Caption Vault styling (title block + Heading 1 sections).
Run AFTER gen_fantasy_pack.py completes.
"""
import csv, os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

PACK = "loras/realistic_posts/fantasy_pack"
MANIFEST = os.path.join(PACK, "fantasy_pack_manifest.csv")
OUT_DOCX = r"C:\Users\David\Documents\Sales\Fantasy_Author_Image_Pack_100.docx"
IMG_W = Inches(3.1)

# brand-ish colors
NAVY = RGBColor(0x1F, 0x2A, 0x44)
GOLD = RGBColor(0xB8, 0x8A, 0x2B)

doc = Document()

# ---- base styles to mimic Caption Vault look ----
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)

def set_cell_text(cell, text, size=9, bold=False, color=None, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align: p.alignment = align
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    if color: run.font.color.rgb = color

def add_bookmark(paragraph, name):
    """Insert a Word bookmark (explicit anchor) into a paragraph."""
    run = paragraph.add_run()
    o = OxmlElement("w:bookmarkStart")
    o.set(qn("w:id"), str(abs(hash(name)) % 1000000))
    o.set(qn("w:name"), name)
    run._r.append(o)
    run2 = paragraph.add_run()
    o2 = OxmlElement("w:bookmarkEnd")
    o2.set(qn("w:id"), str(abs(hash(name)) % 1000000))
    run2._r.append(o2)

def add_hyperlink(paragraph, text, anchor):
    """Internal hyperlink to a bookmark name."""
    # build the hyperlink element
    h = OxmlElement("w:hyperlink")
    h.set(qn("w:anchor"), anchor)
    h.set(qn("w:history"), "1")
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")  # classic link blue
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(color); rPr.append(u)
    t = OxmlElement("w:t")
    t.text = text
    r.append(rPr); r.append(t)
    h.append(r)
    paragraph._p.append(h)

def shade_cell(cell, fill="07111F"):
    """Apply background shading to a table cell (matches other Inkblade products)."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

# ---- Title block ----
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("INKBLADE AUTHOR STUDIO")
r.bold = True; r.font.size = Pt(20); r.font.color.rgb = NAVY
s = doc.add_paragraph()
s.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = s.add_run("Fantasy Author Image Pack")
r.bold = True; r.font.size = Pt(16); r.font.color.rgb = GOLD
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("100 Realistic, Commercial-Use Fantasy Images for Authors, RPG Makers, and Worldbuilders")
r.italic = True; r.font.size = Pt(11)
doc.add_paragraph("")

# ---- Start Here section (matches other Inkblade products) ----
doc.add_heading("Start Here", level=1)
start = doc.add_paragraph(
    "Thank you for purchasing the Fantasy Author Image Pack. This collection of 100 "
    "realistic fantasy images was created to help you promote your books, build ads, "
    "design covers and chapter art, and bring your world to life without hiring an artist "
    "or running your own generators.")
start.style = doc.styles["Body Text"]
use1 = doc.add_paragraph(
    "Each image is royalty-free for commercial use in your own author business: book ads, "
    "social posts, mockups, blog headers, Patreon art, and store visuals. You may edit, "
    "crop, recolor, or combine them however you like.")
use1.style = doc.styles["Body Text"]
use2 = doc.add_paragraph(
    "Use the Image Index below to jump straight to any asset — every caption is a clickable "
    "link that takes you to that image in the gallery. Themes include elves, knights, "
    "dragons, magic, and castles.")
use2.style = doc.styles["Body Text"]

# ---- read manifest ----
with open(MANIFEST, newline="", encoding="utf-8") as f:
    rows = list(csv.DictReader(f))
assert len(rows) == 100, len(rows)

# ---- Reference table section (clickable) ----
doc.add_heading("Image Index (Click to Jump)", level=1)
intro = doc.add_paragraph("Each caption links directly to its image in the gallery below. "
                          "Use this to find the perfect asset fast.")
intro.style = doc.styles["Body Text"]
ref = doc.add_table(rows=1, cols=3)
ref.style = "Table Grid"
ref.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = ref.rows[0].cells
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
for c, txt in zip(hdr, ["#", "Theme", "Image (click to view)"]):
    set_cell_text(c, txt, size=10, bold=True, color=WHITE)
    shade_cell(c, "07111F")
for row in rows:
    idx = int(row["index"]); theme = row["theme"].title(); cap = row["caption"]
    anchor = f"img_{idx:03d}"
    cells = ref.add_row().cells
    set_cell_text(cells[0], str(idx), size=9)
    set_cell_text(cells[1], theme, size=9)
    # caption cell = hyperlink to bookmark
    cells[2].text = ""
    p = cells[2].paragraphs[0]
    add_hyperlink(p, cap, anchor)

doc.add_page_break()

# ---- Gallery: 4 per page (2x2) ----
doc.add_heading("Image Gallery", level=1)
for start in range(0, 100, 4):
    block = rows[start:start+4]
    tbl = doc.add_table(rows=2, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, row in enumerate(block):
        idx = int(row["index"]); fname = row["filename"]; cap = row["caption"]
        cell = tbl.cell(j // 2, j % 2)
        cell.text = ""
        pimg = cell.paragraphs[0]
        pimg.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_bookmark(pimg, f"img_{idx:03d}")
        img_path = os.path.join(PACK, fname)
        if os.path.exists(img_path):
            pimg.add_run().add_picture(img_path, width=IMG_W)
        pcap = cell.add_paragraph()
        pcap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rc = pcap.add_run(f"{idx}. {cap}")
        rc.font.size = Pt(8); rc.italic = True
    doc.add_page_break()

doc.save(OUT_DOCX)
print("SAVED:", OUT_DOCX)
print("reference rows:", len(ref.rows)-1, "| gallery pages:", 100//4)
